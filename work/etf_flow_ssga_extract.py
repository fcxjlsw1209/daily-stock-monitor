"""Read-only issuer workbook extraction; no flow, score, order or return math."""
from datetime import date, datetime, timezone
from pathlib import Path
import hashlib
import json
import math
from openpyxl import load_workbook

ROOT = Path('/Users/shiweili/code/monthly-stock-research')
OUT = ROOT / 'outputs/monthly_etf_primary_flow_sources/ssga'
SYMBOLS = ('SPY', 'GLD', 'XLE', 'XLF', 'XLI', 'XLK', 'XLV', 'XLY')
HEADERS = ['Date', 'NAV', 'Shares Outstanding', 'Total Net Assets']

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def positive(value):
    return (isinstance(value, (int, float)) and not isinstance(value, bool)
            and math.isfinite(value) and value > 0)

def date_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return datetime.strptime(str(value), '%d-%b-%Y').date().isoformat()

def dump_new(path, payload):
    with path.open('x') as handle:
        json.dump(payload, handle, indent=2, ensure_ascii=False, allow_nan=False)
        handle.write('\n')

def main():
    targets = [OUT / (s + '_SOURCE_EXTRACT.json') for s in SYMBOLS]
    targets.append(OUT / 'FEASIBILITY.json')
    if any(p.exists() for p in targets):
        raise RuntimeError('Write-once source extraction already exists; inspect instead of overwriting')
    summary = []
    for symbol in SYMBOLS:
        raw = OUT / (symbol + '_navhistory.xlsx')
        receipt_path = OUT / (symbol + '_navhistory_RETRIEVAL.json')
        receipt = json.loads(receipt_path.read_text())
        assert receipt['status'] == 200
        assert digest(raw) == receipt['sha256']
        workbook = load_workbook(raw, read_only=True, data_only=True)
        sheet = workbook['navhist']
        header_rows = list(sheet.iter_rows(min_row=1, max_row=4, values_only=True))
        assert str(header_rows[1][1]).replace('®', '').strip() == symbol
        assert list(header_rows[3][:4]) == HEADERS
        rows, notes = [], []
        for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
            if number <= 4 or not any(v is not None for v in values):
                continue
            try:
                day = date_value(values[0])
            except (ValueError, TypeError):
                notes.append({'source_row': number, 'nonempty_cells': [str(v) for v in values if v is not None]})
                continue
            rows.append({'date': day, 'source_row': number, 'raw_date': str(values[0]),
                         'nav': values[1], 'shares_outstanding': values[2], 'net_assets': values[3]})
        physical_rows, physical_cols = sheet.max_row, sheet.max_column
        workbook.close()
        days = [r['date'] for r in rows]
        assert len(set(days)) == len(days), (symbol, 'duplicate observation dates')
        research = [r for r in rows if '2016-08-17' <= r['date'] <= '2026-09-24']
        fields = {}
        for field in ('nav', 'shares_outstanding', 'net_assets'):
            good = [r['date'] for r in rows if positive(r[field])]
            fields[field] = {
                'positive_numeric_rows': len(good),
                'first_positive_numeric_date': min(good) if good else None,
                'last_positive_numeric_date': max(good) if good else None,
                'invalid_or_placeholder_rows': len(rows) - len(good),
                'invalid_research_dates': [r['date'] for r in research if not positive(r[field])],
            }
        entry = {
            'symbol': symbol, 'fund_name': header_rows[0][1], 'sheet': 'navhist',
            'header_range': 'A4:D4', 'source_file': str(raw.relative_to(ROOT)),
            'source_sha256': digest(raw), 'receipt_sha256': digest(receipt_path),
            'source_url': receipt['requested_url'], 'retrieved_at_utc': receipt['retrieved_utc'],
            'date_rows': len(rows), 'first_date': min(days), 'last_date': max(days),
            'duplicate_dates': 0, 'physical_shape': [physical_rows, physical_cols],
            'weekend_dates': [d for d in days if date.fromisoformat(d).weekday() >= 5],
            'research_date_rows': len(research), 'fields': fields,
            'session_calendar_audit': 'NOT_YET_RUN',
            'historical_publication_clock': 'UNKNOWN',
            'share_trade_or_settlement_basis': 'UNKNOWN',
            'split_restatement_convention': 'UNKNOWN_PENDING_SEPARATE_OFFICIAL_ACTION_AUDIT',
            'strict_point_in_time_admission': False,
        }
        extraction = {
            'scope': 'Original-source field extraction and coverage only. Dates are not claimed to be sessions or first-public timestamps.',
            'created_at_utc': datetime.now(timezone.utc).isoformat(),
            'parser_sha256': digest(Path(__file__)), 'metadata': entry,
            'headers': HEADERS, 'rows': rows, 'source_notes': notes,
            'raw_values_preserved': True, 'flow_or_strategy_calculations': False,
        }
        destination = OUT / (symbol + '_SOURCE_EXTRACT.json')
        dump_new(destination, extraction)
        entry['extraction_file'] = str(destination.relative_to(ROOT))
        entry['extraction_sha256'] = digest(destination)
        summary.append(entry)
        print(json.dumps({'symbol': symbol, 'dates': len(rows), 'first': min(days),
                          'first_shares': fields['shares_outstanding']['first_positive_numeric_date'],
                          'research_missing_shares': len(fields['shares_outstanding']['invalid_research_dates'])}))
    dump_new(OUT / 'FEASIBILITY.json', {
        'created_at_utc': datetime.now(timezone.utc).isoformat(),
        'status': 'OFFICIAL_RAW_VALUES_COLLECTED_NOT_STRICT_PIT_ADMITTED',
        'sources': summary, 'source_workbooks_modified': False,
        'new_flow_or_strategy_calculations': 0, 'observer_changes': False,
        'limitations': ['Publication and historical revision clocks not established.',
                        'Split and action conventions require a separate audit.',
                        'No missing row or placeholder is filled, converted to zero or silently discarded.',
                        '2016 research coverage counts are source dates, not verified trading sessions.']})

if __name__ == '__main__':
    main()
